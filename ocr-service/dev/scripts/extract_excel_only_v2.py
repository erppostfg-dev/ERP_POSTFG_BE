"""
PO extraction for channels that send the whole PO as one spreadsheet (xls/xlsx/csv),
e.g. Flipkart quick-commerce. No PDF and no OCR: line items and totals are read straight from the
cells. Only the header block above the table goes to the Bedrock model, which pairs
labels with values. The header and line items are then checked with plain rules
(GSTIN/PAN format, row arithmetic, footer totals vs column sums).

v2: temperature 0 on the model call so field names stay stable between runs.

Usage:
    python extract_excel_only_v2.py <PO file: s3://bucket/key or local>

Prints the extracted PO as JSON; the log goes to stderr.

Needs: pip install boto3 openpyxl xlrd, and AWS credentials (aws configure).
"""

import os
import sys
import re
import io
import time
import json
import csv
import contextlib
import boto3
import openpyxl
from botocore.config import Config
from botocore.exceptions import ClientError, NoCredentialsError

AWS_REGION = "ap-south-1"           # S3 region

BEDROCK_REGION = "ap-south-1"
# Qwen3 VL runs in ap-south-1, so the PO data stays in Mumbai. Text-only call here.
BEDROCK_MODEL_ID = "qwen.qwen3-vl-235b-a22b"
BEDROCK_MAX_TOKENS = 4000           # header JSON is small
TEMPERATURE = 0.0                   # keeps field names the same between runs

# Mumbai on-demand pricing, only used for the cost printout. Check the Bedrock pricing page now and then.
MODEL_INPUT_COST_PER_M = 0.62       # USD per 1M input tokens
MODEL_OUTPUT_COST_PER_M = 3.13      # USD per 1M output tokens

ROUNDING_TOLERANCE = 1.0            # rupees per row, some POs round to whole rupees
MAX_FAILING_SHARE = 0.5             # if more rows than this fail, the formula doesn't fit the layout


def make_clients() -> dict:

    config = Config(retries={"max_attempts": 10, "mode": "adaptive"}, read_timeout=600)
    return {
        "s3": boto3.client("s3", region_name=AWS_REGION, config=config),
        "bedrock": boto3.client("bedrock-runtime", region_name=BEDROCK_REGION, config=config),
    }


def parse_s3_uri(s3_uri: str) -> tuple:
    without_prefix = s3_uri[len("s3://"):]
    bucket, _, key = without_prefix.partition("/")
    if not bucket or not key:
        raise ValueError(f"Could not parse bucket/key from S3 path: {s3_uri}")
    return bucket, key


def read_file_bytes(s3, path: str) -> bytes:
    """Read an s3:// path or a local file."""
    if path.startswith("s3://"):
        bucket, key = parse_s3_uri(path)
        return s3.get_object(Bucket=bucket, Key=key)["Body"].read()
    with open(path, "rb") as f:
        return f.read()


# --- reading the sheet ---

XLSX_SIGNATURE = b"PK\x03\x04"                            # xlsx/xlsm are zip files
XLS_SIGNATURE = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"       # old binary Excel (OLE)
TEXT_EXTENSIONS = ("csv", "tsv", "txt")
SUPPORTED_EXTENSIONS = ("xlsx", "xlsm", "xls") + TEXT_EXTENSIONS


def detect_format(path: str, data: bytes) -> str:
    """Work out xlsx/xls/tsv/csv from the file signature, falling back to the extension."""
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    if ext and ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported PO file type '.{ext}' "
                         f"(supported: {', '.join('.' + e for e in SUPPORTED_EXTENSIONS)})")

    if data[:4] == XLSX_SIGNATURE:
        fmt = "xlsx"
    elif data[:8] == XLS_SIGNATURE:
        fmt = "xls"
    elif ext == "tsv":
        fmt = "tsv"
    else:
        fmt = "csv"  # also covers a .xls that is really a text export

    if ext and ext not in (fmt, "xlsm", "txt"):
        print(f"  [Input] extension is .{ext} but content looks like {fmt.upper()}, reading as {fmt.upper()}")
    return fmt


def cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def xlsx_to_rows(data: bytes) -> list:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        # Scootsy's export claims the sheet is "A1:A1" and read-only mode believes it.
        ws.reset_dimensions()
        return [[cell_to_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def xls_to_rows(data: bytes) -> list:
    """Date cells come out as YYYY-MM-DD rather than Excel serials."""
    try:
        import xlrd  # only needed for old-style .xls
    except ImportError:
        raise RuntimeError("xlrd is needed for .xls files: pip install xlrd")
    xlrd_log = io.StringIO()  # keep xlrd's warnings off the console
    try:
        book = xlrd.open_workbook(file_contents=data, logfile=xlrd_log)
    except xlrd.compdoc.CompDocError as e:
        # FirstCry's xls has a slightly broken container (a sector listed twice). Excel
        # opens it and the data is fine, so retry in tolerant mode.
        print(f"  [Input] .xls container is slightly malformed ({e}), "
              f"retrying in tolerant mode")
        book = xlrd.open_workbook(file_contents=data, logfile=xlrd_log, ignore_workbook_corruption=True)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode).strftime("%Y-%m-%d"))
            else:
                row.append(cell_to_str(cell.value))
        rows.append(row)
    return rows


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding) 
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def text_to_rows(data: bytes, fmt: str) -> list:
    text = decode_text(data)
    if fmt == "tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_sheet_rows(path: str, data: bytes) -> tuple:
    fmt = detect_format(path, data)
    if fmt == "xls":
        rows = xls_to_rows(data)
    elif fmt == "xlsx":
        rows = xlsx_to_rows(data)
    else:
        rows = text_to_rows(data, fmt)
    # openpyxl dates come out as "2026-10-06 00:00:00"; drop the midnight time.
    rows = [[re.sub(r"^(\d{4}-\d{2}-\d{2}) 00:00:00$", r"\1", v) for v in row] for row in rows]
    return fmt, rows


# --- line items and totals ---

CURRENCY_TOKENS = re.compile(r"\bINR\b|₹|\bRs\.?", re.IGNORECASE)
SERIAL_HEADER_RE = re.compile(r"^((s|sr|sl)\.?\s*no\.?|#|no\.?|line\s*no\.?|item\s*no\.?)$", re.IGNORECASE)
QTY_HEADER_RE = re.compile(r"^((order(ed)?|po|billed|free|total|invoice(d)?|dispatch(ed)?|case)\s+)?"
                           r"(qty|quantity)\.?(\s*\(.*\))?$", re.IGNORECASE)
MIN_HEADER_CELLS = 5                # min column names for a header row with no serial column


def to_num(value):
    """Parse '7611.84 INR', '1,234' or '18.0%' as a float. None if it isn't a number."""
    text = CURRENCY_TOKENS.sub("", str(value)).replace(",", "").replace("%", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def format_num(value: float) -> str:
    """Short string for a float, e.g. 55902.240000000005 gives '55902.24', 499.0 gives '499'."""
    text = f"{round(value, 4):.4f}".rstrip("0").rstrip(".")
    return "0" if text == "-0" else text


def clean_cell(value: str) -> str:
    """Collapse whitespace and strip currency from amounts. Plain codes are left alone
    so leading zeros in HSN/EAN survive."""
    value = " ".join(value.split())
    if CURRENCY_TOKENS.search(value):
        num = to_num(value)
        if num is not None:
            return format_num(num)
    return value


def find_table_header(rows: list) -> tuple:
    """
    Find the line-items header row: the first row starting with a serial column
    ("S. no.", "#", ...), otherwise the first all-text row with a Quantity column
    (Myntra has no serial column). Returns (header_row_index, serial_col, qty_col).
    """
    def qty_col(row):
        return next((j for j, c in enumerate(row) if QTY_HEADER_RE.match(" ".join(c.split()))), None)

    for i, row in enumerate(rows):
        if row and SERIAL_HEADER_RE.match(row[0].strip()):
            if qty_col(row) is None:
                # Second line of a two-row header (Scootsy). Qty is in the row above.
                prev = next((p for p in range(i - 1, -1, -1) if any(v.strip() for v in rows[p])), None)
                if prev is not None and qty_col(rows[prev]) is not None:
                    return prev, 0, qty_col(rows[prev])
            return i, 0, qty_col(row)
    for i, row in enumerate(rows):
        filled = [c for c in row if c.strip()]
        if len(filled) >= MIN_HEADER_CELLS and not any(to_num(c) is not None for c in filled) \
                and qty_col(row) is not None:
            return i, None, qty_col(row)
    raise RuntimeError("Could not find the line-items table: no 'S. no.' or '#' column, and no header "
                       "row with a Quantity column.")


ITEM_COUNT_RE = re.compile(r"^\(\s*count\s*\)\s*(\d+)$", re.IGNORECASE)     # e.g. BigBasket "(count) 1"


def footer_value(value) -> bool:
    """True for a number or an item count like "(count) 1"."""
    return to_num(value) is not None or bool(ITEM_COUNT_RE.match(str(value).strip()))


def read_footer_row(row: list, header: list, col_sums: dict = None, rate_cols: set = None,
                    n_items: int = 0) -> list:
    """
    Pull the totals out of one footer row: label/value pairs, item counts, a row title
    with numbers under their columns, and bare sums under a column. Each total's
    "target" is the column to reconcile against, "count", or None. col_sums lets a sum
    that sits in the wrong column (Scootsy) be matched to a nearby column it equals.
    """
    col_sums, rate_cols = col_sums or {}, rate_cols or set()

    def header_at(j):
        return header[j] if j < len(header) else ""

    found, used = [], set()
    for j, cell in enumerate(row):
        label = cell.strip()
        if j in used or not label or footer_value(label):
            continue
        k = next((k for k in range(j + 1, len(row)) if row[k].strip()), None)
        if k is None or not footer_value(row[k]):
            continue
        used.update({j, k})
        count = ITEM_COUNT_RE.match(row[k].strip())
        if count:
            found.append((k, {"field": "Total Items", "value": count.group(1), "target": "count"}))
            continue
        if j == 0 and k > j + 1 and header_at(k):
            # Row title like FirstCry's "Sub Total", with numbers further right under
            # their columns. Name each number after its column.
            title = label.rstrip("=:").strip()
            # Net/discounted totals won't match a plain column sum, so don't reconcile them.
            adjusted = re.search(r"discount|\bnet\b|round|payable|after", title, re.IGNORECASE)
            used.discard(k)                         # k was reserved above as a pair value
            for c in range(k, len(row)):
                if c not in used and to_num(row[c]) is not None:
                    used.add(c)
                    col_name = header_at(c) or f"Column {c + 1}"
                    found.append((c, {"field": f"{title} ({col_name})", "value": format_num(to_num(row[c])),
                                      "target": None if adjusted else c}))
            continue
        name = label.rstrip("=:").strip()
        target = None
        if name.lower() == "total" and header_at(k):
            name = f"Total {header_at(k)}"          # bare "Total=": name it after its column
            target = k
        core = re.sub(r"^total\s*", "", name, flags=re.IGNORECASE).lower()
        if target is None:          # column named like the label, or like it minus "Total"
            target = next((i for i, h in enumerate(header) if h.lower() in (name.lower(), core)), None)
        if target is None and re.fullmatch(r"(items|lines|line items|no\.? of items|skus)", core):
            target = "count"
        found.append((k, {"field": name, "value": format_num(to_num(row[k])), "target": target}))

    # Unlabelled numbers usually sit under the column they total. If not, look up to 3
    # columns either side for one whose sum matches (skipping rate/% columns). Failing
    # that, keep it under its own column, or unreconciled if there's nothing to check.
    claimed = {t["target"] for _, t in found if isinstance(t["target"], int)}
    tolerance = max(0.05, 0.01 * n_items)

    def total_name(c):
        col_name = header_at(c) or f"Column {c + 1}"
        return col_name if col_name.lower().startswith("total") else f"Total {col_name}"

    for j, cell in enumerate(row):
        if j in used or to_num(cell) is None:
            continue
        value = to_num(cell)
        direct = j in col_sums and j not in claimed and j not in rate_cols
        if direct and abs(col_sums[j] - value) <= tolerance:
            target = j
        else:
            target = next((c for c in (j + 1, j - 1, j + 2, j - 2, j + 3, j - 3)
                           if c in col_sums and c not in claimed and c not in rate_cols
                           and abs(col_sums[c] - value) <= tolerance), None)
            if target is None and direct:
                target = j
        if target is None:
            found.append((j, {"field": f"Footer value (column {j + 1})", "value": format_num(value), "target": None}))
            continue
        claimed.add(target)
        found.append((j, {"field": total_name(target), "value": format_num(value), "target": target}))

    return [total for _, total in sorted(found, key=lambda t: t[0])]


def read_line_items(rows: list, header_idx: int, serial_col, qty_col) -> dict:
    """
    Item rows are the ones after the header with a number in the serial (or Qty) column.
    Rows after that which still hold numbers are footer totals; the first row with no
    numbers (T&C etc.) ends the table.
    """
    header = [" ".join(v.split()) for v in rows[header_idx]]
    first_body_row = header_idx + 1

    # Two-row header (Scootsy: "CGST" over "Rate" | "Amt (INR)"): merge into one name per column.
    nxt = next((i for i in range(header_idx + 1, len(rows)) if any(v.strip() for v in rows[i])), None)
    if nxt is not None and not any(footer_value(v) for v in rows[nxt] if v.strip()):
        sub = [" ".join(v.split()) for v in rows[nxt]]
        width = max(len(header), len(sub))
        header, sub = header + [""] * (width - len(header)), sub + [""] * (width - len(sub))
        merged = []
        for j in range(width):
            top = header[j]
            if not top and sub[j]:          # merged heading cell, take it from the left
                top = next((header[k] for k in range(j - 1, -1, -1) if header[k]), "")
            merged.append(f"{top} {sub[j]}".strip())
        header = merged
        first_body_row = nxt + 1
        print(f"  [Line items] merged two-row header: {', '.join(h for h in merged if ' ' in h)[:120]}")

    width = len(header)
    key_col = serial_col if serial_col is not None else qty_col

    raw_rows, footer_rows, in_footer = [], [], False
    for row in rows[first_body_row:]:
        if not any(v.strip() for v in row):
            if raw_rows:
                in_footer = True                    # blank line ends the items
            continue
        padded = (row + [""] * width)[:width]
        if not in_footer and key_col is not None and to_num(padded[key_col]) is not None:
            raw_rows.append(padded)
            continue
        in_footer = True
        if any(footer_value(v) for v in row if v.strip()):
            footer_rows.append(row)
        else:
            break

    # Drop fully empty columns (e.g. leftovers of a merged cell) and name any unnamed ones.
    keep = [i for i in range(width) if header[i] or any(r[i].strip() for r in raw_rows)]
    columns = [header[i] or f"Column {i + 1}" for i in keep]
    items = [[clean_cell(r[i]) for i in keep] for r in raw_rows]

    currency = "INR" if any(CURRENCY_TOKENS.search(v) for r in raw_rows for v in r) else ""
    col_sums = {}
    for i in range(width):
        values = [to_num(r[i]) for r in raw_rows]
        if any(v is not None for v in values):
            col_sums[i] = sum(v for v in values if v is not None)
    rate_cols = {i for i, h in enumerate(header) if re.search(r"rate|%", h, re.IGNORECASE)}
    totals = [t for row in footer_rows
              for t in read_footer_row(row, header, col_sums, rate_cols, len(raw_rows))]

    return {"header": header, "raw_rows": raw_rows, "columns": columns, "rows": items,
            "totals": totals, "currency": currency}


def reconcile_totals(header: list, raw_rows: list, totals: list) -> tuple:
    """Check footer totals against column sums (or the item count). Returns (warnings, summary)."""
    warnings, checked = [], []
    for total in totals:
        target, expected = total["target"], to_num(total["value"])
        if target is None or expected is None:
            continue
        if target == "count":
            checked.append(total["field"])
            if int(expected) != len(raw_rows):
                warnings.append(f"Footer '{total['field']}' = {total['value']} but the sheet has "
                                f"{len(raw_rows)} line item(s)")
            continue
        values = [v for v in (to_num(r[target]) for r in raw_rows if target < len(r)) if v is not None]
        if not values:
            continue
        actual = sum(values)
        if QTY_HEADER_RE.match(header[target]):
            tolerance = 0.001
        elif all(v == int(v) for v in values):
            tolerance = max(ROUNDING_TOLERANCE, 0.5 * ROUNDING_TOLERANCE * len(values))
        else:
            tolerance = max(0.05, 0.01 * len(values))
        checked.append(total["field"])
        if abs(expected - actual) > tolerance:
            warnings.append(f"Footer '{total['field']}' = {total['value']} but the "
                            f"'{header[target]}' column sums to {format_num(actual)}")
    if not checked:
        return warnings, "totals check: skipped (no footer total matched a line-item column)"
    return warnings, f"totals check: {', '.join(checked)} vs the line items"


# Column names (normalised) the row arithmetic check looks for. Billed Qty comes before
# Total Qty because Total Qty can include free units.
QTY_COLUMNS = ("qty", "qty.", "quantity", "billed qty", "billed quantity", "order qty", "ordered qty",
               "po qty", "invoice qty", "total qty")
RATE_COLUMNS = ("supplier price", "rate/unit", "rate", "unit rate", "unit price", "landing rate",
                "landing cost", "landed price", "landed cost", "basic cost price", "basic cost",
                "base cost", "unit base cost", "unit cost", "cost price", "list price", "price")
AMOUNT_COLUMNS = ("total amount", "total value", "amount (inr)", "amount", "total amt", "total amt.",
                  "net amount", "line total", "total")
TAXABLE_COLUMNS = ("taxable value", "taxable amount", "taxable amt")
# Tax value columns, including a bare "IGST" (BigBasket, next to "IGST%"). A combined
# "GST Amount" is left out on purpose since it would double-count IGST or CGST+SGST.
TAX_VALUE_COLUMN_RE = re.compile(r"^(igst|cgst|sgst|utgst|cess)\b.*\b(value|amt\.?|amount)$"
                                 r"|^tax (amt\.?|amount)$|^(igst|cgst|sgst|utgst)$")
TAX_PERCENT_COLUMN_RE = re.compile(r"^(igst|cgst|sgst|utgst|cess)\b.*(%|rate)$")


def _norm_col(name: str) -> str:
    """Lower-case, collapse spaces and drop a "(INR)" style suffix."""
    name = re.sub(r"\s*\((inr|rs\.?|₹)\)\s*", " ", str(name), flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", name.strip().lower())


def validate_line_items(columns: list, rows: list) -> tuple:
    """
    Per-row arithmetic, run as two separate checks:
      rate:    amount = Qty x Rate, plus tax (as values or %), or with a tax-inclusive rate
      taxable: amount = Taxable Value plus tax
    If most rows fail a check, it's reported once as not applicable to this layout.
    """
    names = [_norm_col(c) for c in columns]
    qty_idx = next((names.index(c) for c in QTY_COLUMNS if c in names), None)
    rate_idxs = [names.index(c) for c in RATE_COLUMNS if c in names]
    amount_idxs = [names.index(c) for c in AMOUNT_COLUMNS if c in names]
    taxable_idx = next((names.index(c) for c in TAXABLE_COLUMNS if c in names), None)
    tax_value_idxs = [i for i, n in enumerate(names) if TAX_VALUE_COLUMN_RE.search(n)]
    tax_pct_idxs = [i for i, n in enumerate(names) if TAX_PERCENT_COLUMN_RE.search(n)]
    if qty_idx is None or not amount_idxs or not (rate_idxs or taxable_idx is not None):
        return [], "row arithmetic check: skipped (no Qty / Rate / Amount columns found by name)"

    def cell(row, i):
        return to_num(row[i]) if i is not None and i < len(row) else None

    def fits(amounts, candidates):
        return any(abs(a - c) <= max(ROUNDING_TOLERANCE, abs(a) * 0.001) for a in amounts for c in candidates)

    def decimal_hint(amounts, candidates):
        return (" (looks like a misplaced decimal point)"
                if any(c and abs(a / c - f) < 0.02 * f
                       for a in amounts for c in candidates for f in (10, 100, 0.1, 0.01)) else "")

    checks = {"rate": {"checked": 0, "failures": []}, "taxable": {"checked": 0, "failures": []}}
    for r, row in enumerate(rows, start=1):
        qty = cell(row, qty_idx)
        amounts = [a for a in (cell(row, ai) for ai in amount_idxs) if a is not None]
        amount_desc = f"{columns[amount_idxs[0]]} {row[amount_idxs[0]]}" if amounts else ""
        if not amounts and cell(row, taxable_idx) is not None:
            # No line total on the row (Scootsy): compare Qty x Rate with Taxable Value and
            # skip the taxable check.
            amounts = [cell(row, taxable_idx)]
            amount_desc = f"{columns[taxable_idx]} {row[taxable_idx]}"
            taxable_only = True
        else:
            taxable_only = False
        if not qty or not amounts:
            continue
        tax_values = sum(v for v in (cell(row, i) for i in tax_value_idxs) if v is not None)
        tax_pct = sum(v for v in (cell(row, i) for i in tax_pct_idxs) if v is not None)

        rate_candidates = []
        for ri in rate_idxs:
            rate = cell(row, ri)
            if rate is not None:
                base = qty * rate
                rate_candidates += [base, base + tax_values, base * (1 + tax_pct / 100)]
        if rate_candidates:
            checks["rate"]["checked"] += 1
            if not fits(amounts, rate_candidates):
                checks["rate"]["failures"].append(
                    f"Row {r}: Qty({row[qty_idx]}) x {columns[rate_idxs[0]]}({row[rate_idxs[0]]}) doesn't give "
                    f"{amount_desc} with or without tax{decimal_hint(amounts, rate_candidates)}")

        taxable = cell(row, taxable_idx)
        if taxable is not None and not taxable_only:
            taxable_candidates = [taxable + tax_values, taxable * (1 + tax_pct / 100)]
            checks["taxable"]["checked"] += 1
            if not fits(amounts, taxable_candidates):
                checks["taxable"]["failures"].append(
                    f"Row {r}: {columns[taxable_idx]}({row[taxable_idx]}) + tax doesn't give "
                    f"{amount_desc}{decimal_hint(amounts, taxable_candidates)}")

    warnings, notes = [], []
    for label, c in (("Qty x Rate", checks["rate"]), ("Taxable + tax", checks["taxable"])):
        if not c["checked"]:
            continue
        if len(c["failures"]) > MAX_FAILING_SHARE * c["checked"] and c["checked"] >= 3:
            notes.append(f"{label}: not applied ({len(c['failures'])} of {c['checked']} rows fit no formula, "
                         f"this layout computes amounts differently)")
        else:
            warnings += c["failures"]
            notes.append(f"{label}: {c['checked'] - len(c['failures'])} of {c['checked']} rows OK")
    if not notes:
        return [], "row arithmetic check: skipped (no numeric rows)"
    return warnings, "row arithmetic check: " + "; ".join(notes)


# --- header fields (model) ---

def _cell_lines(value: str) -> list:
    return [" ".join(line.split()) for line in str(value).split("\n") if line.strip()]


def _is_stacked_labels(lines: list) -> bool:
    """Several short labels stacked in one cell, one per line."""
    return len(lines) > 1 and all(len(line) <= 40 and not any(ch.isdigit() for ch in line) for line in lines)


def header_section_text(rows: list, header_idx: int) -> str:
    """
    Rows above the table as "label | value | ..." lines for the prompt. A cell of stacked
    labels next to a cell with the same number of values (FirstCry) is split into pairs.
    """
    out = []
    for row in rows[:header_idx]:
        raw = [v for v in row if v.strip()]
        cells, i = [], 0
        while i < len(raw):
            labels = _cell_lines(raw[i])
            values = _cell_lines(raw[i + 1]) if i + 1 < len(raw) else []
            if _is_stacked_labels(labels) and len(values) == len(labels):
                for label, value in zip(labels, values):
                    cells += [label, value]
                i += 2
                continue
            cells.append(" ".join(raw[i].split()))
            i += 1
        # Mark a trailing "Label:" as blank, otherwise the model tends to glue the next
        # row onto it (Scootsy's "Reference PO Code:" got the vendor block).
        if cells and cells[-1].rstrip().endswith(":"):
            cells.append(BLANK_MARKER)
        if cells:
            out.append(" | ".join(cells))
    return "\n".join(out)


BLANK_MARKER = "(blank)"


# Plain string (not .format()-ed) so the JSON example keeps single braces.
HEADER_FIELDS_PROMPT = """You are extracting the HEADER fields of a Purchase Order (or Stock Transfer) \
that was exported as a spreadsheet. Below is every non-empty row ABOVE the line-items table. Within \
a row, the non-empty cells are separated by " | " in left-to-right order. These are the file's \
exact cell values — copy them character for character.

HOW THE SHEET IS LAID OUT
- A label is usually followed by its value in the next cell (e.g. "PO# | FASWN08497679"), and one \
row can hold several label/value pairs side by side (e.g. "PO Expiry | 06-10-26 | CATEGORY | \
Grooming | ORDER DATE | 17-09-26").
- One cell can hold a whole BLOCK of text — e.g. a "Registered Address" label followed by one cell \
containing the company name, address, phone, fax and CIN. Return the block as its own field (e.g. \
"Registered Address"), AND return every labelled value inside it as a separate field too (e.g. \
"Registered Office CIN", "Registered Office Tel", "Registered Office Fax").
- A title cell can carry a value (e.g. "PURCHASE ORDER # FASWN08497679"); use it if the same value \
isn't labelled elsewhere.
- An address block often starts with a line that is only the location's NAME (e.g. "Warehouse \
Address" followed by "WMS-Kolkata-FC", then the street address lines). Return that first line as \
"<Block> Name" (e.g. "Warehouse Name", "Delivery Name") and the remaining lines as "<Block> \
Address" — every field name must be unique, never return the same name twice.
- A value cell may start with a leftover label fragment (e.g. "GSTIN NO:GST: 19AACCI2053A1Z4"); \
return only the value itself ("19AACCI2053A1Z4").
- "(blank)" after a label means NOTHING is printed after that label on its row — the label has no \
value. Omit it, and never attach the text of the next row (e.g. an address or company block) to it.
- A block of text with no label of its own (e.g. one cell holding a company name, address, GSTIN, \
PAN and contact) belongs to the nearest heading above it that has no value on its own row (e.g. \
"Vendor Name :"). Name its fields after that heading (e.g. "Vendor Address", "Vendor GSTIN", \
"Vendor PAN", "Vendor Contact").

WHAT TO EXTRACT
Every labelled business field: buyer / supplier names, PO number (for stock transfers: STR / STN / \
STO number), dates (order date, expiry, delivery), addresses, GSTIN / PAN / CIN, state codes, \
contact numbers / emails, payment mode, credit / payment terms, contract references, nature of \
supply / transaction, category, etc.

RULES
- Copy each value EXACTLY as it appears — do not reformat dates, numbers or codes, and do not \
shorten, de-duplicate or rewrite addresses (even if the address text repeats itself).
- When the same label appears more than once (e.g. GSTIN / State Code next to the billing, \
shipping and supplier addresses), name each one after the block it belongs to (e.g. "<Block> \
GSTIN", "<Block> State Code") — return every occurrence, never merge them.
- Name a block after its heading AS PRINTED on the sheet — "Billing Address" gives "Billing Name", \
"Billing GSTIN"; "Shipping Address" gives "Shipping …". Never reinterpret a heading into a \
different party or direction (a "Billing Address" block is not "Billed By").
- Every GSTIN, PAN and CIN in the rows must be returned as its own field.
- Skip section titles that have no value of their own (e.g. "PAYMENT DETAILS") and labels whose \
value is blank or a placeholder such as "_", "-", "NA" — omit them entirely.
- Do NOT extract line items, totals, terms & conditions or notifications.

For EACH field, also self-report your confidence as "high", "medium" or "low" — how certain you \
are that the label/value pairing is right.

Return ONLY valid JSON (no markdown fences, no commentary) in exactly this shape:
{"header_fields": [{"field": "<field name in Title Case>", "value": "<value exactly as in the cell>", "confidence": "high|medium|low"}]}

--- SPREADSHEET ROWS ABOVE THE LINE-ITEMS TABLE ---
"""

PLACEHOLDER_VALUES = {"_", "-", "--", "na", "n/a", "nil", "none", "null", "(blank)", "blank"}


def _same_text(a, b) -> bool:
    """Compare ignoring case, spaces and punctuation."""
    norm = lambda s: re.sub(r"[^a-z0-9]", "", str(s).lower())
    return bool(norm(a)) and norm(a) == norm(b)


def call_model_for_header(bedrock, sheet_text: str) -> dict:
    response = bedrock.converse(
        modelId=BEDROCK_MODEL_ID,
        messages=[{"role": "user", "content": [{"text": HEADER_FIELDS_PROMPT + sheet_text + "\n"}]}],
        inferenceConfig={"maxTokens": BEDROCK_MAX_TOKENS, "temperature": TEMPERATURE},
    )

    output_blocks = response["output"]["message"]["content"]
    response_text = "".join(b.get("text", "") for b in output_blocks).strip()

    usage = response.get("usage", {})
    input_tokens = usage.get("inputTokens", 0)
    output_tokens = usage.get("outputTokens", 0)
    cost = (input_tokens / 1_000_000 * MODEL_INPUT_COST_PER_M) + \
           (output_tokens / 1_000_000 * MODEL_OUTPUT_COST_PER_M)
    print(f"  [Header | {BEDROCK_MODEL_ID}] tokens: "
          f"{input_tokens} in, {output_tokens} out, {input_tokens + output_tokens} total, "
          f"est. cost ${cost:.4f}")

    if response_text.startswith("```"):
        response_text = response_text.strip("`")
        if response_text.lower().startswith("json"):
            response_text = response_text[4:]
        response_text = response_text.strip()

    try:
        parsed = json.loads(response_text)
    except json.JSONDecodeError as e:
        stop_reason = response.get("stopReason", "unknown")
        raise RuntimeError(f"Model did not return valid JSON (stopReason={stop_reason}): {e}\n"
                           f"Raw response:\n{response_text}")

    header_fields = parsed.get("header_fields", [])
    # Multi-line values (addresses) become one line so each fits on one CSV row.
    for field in header_fields:
        field["value"] = " ".join(str(field.get("value", "")).split())
    # Drop echoed headings where the value is just the field name (Scootsy's "Purchase Order").
    echoed = [f.get("field", "") for f in header_fields if _same_text(f.get("field", ""), f["value"])]
    header_fields = [f for f in header_fields if not _same_text(f.get("field", ""), f["value"])]
    if echoed:
        print(f"  [Header] dropped {len(echoed)} field(s) whose value equals the name: {', '.join(echoed)}")
    # Drop blank or placeholder values ("-", "NA", ...).
    blank = [f.get("field", "") for f in header_fields
             if str(f.get("value", "")).strip().lower() in PLACEHOLDER_VALUES | {""}]
    header_fields = [f for f in header_fields
                     if str(f.get("value", "")).strip().lower() not in PLACEHOLDER_VALUES | {""}]
    if blank:
        print(f"  [Header] dropped {len(blank)} blank field(s): {', '.join(blank)}")

    # Names must be unique for anything keyed by field name, so number repeats.
    seen = {}
    for field in header_fields:
        name = str(field.get("field", ""))
        seen[name] = seen.get(name, 0) + 1
        if seen[name] > 1:
            field["field"] = f"{name} {seen[name]}"
            print(f"  [Header] duplicate name '{name}', renamed to '{field['field']}'")
    return {"header_fields": header_fields, "cost": cost}


# --- header validation ---

PAN_RE = re.compile(r"^[A-Z]{5}[0-9]{4}[A-Z]$")
GSTIN_RE = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]$")
GSTIN_CHARSET = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
# IDs in the header rows; each one should come back as its own field.
IDS_IN_TEXT = {
    "GSTIN": re.compile(r"\b[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b"),
    "PAN": re.compile(r"\b[A-Z]{5}[0-9]{4}[A-Z]\b"),
    "CIN": re.compile(r"\b[LU][0-9]{5}[A-Z]{2}[0-9]{4}[A-Z]{3}[0-9]{6}\b"),
}

# Matched against the lower-cased field name.
IS_PAN = re.compile(r"\bpan\b")                                   # not "company"
IS_GSTIN = re.compile(r"\bgst(in)?\b")
# PO number, or STR/STN/STO number for stock transfers.
IS_PO_NUMBER = re.compile(
    r"\b(p\.?\s?o\.?|purchase order|order)\s*(number|no\.?|#|id)?\s*$|"
    r"\b(p\.?\s?o\.?|purchase order)\s*(number|no\.?|#)|"
    r"\b(str|stn|sto|stock transfer(\s+(request|note|order))?)\s*(number|no\.?|#|id)(?!\w)"
)
IS_ID_FIELD = re.compile(r"\b(number|no\.?|code|id|ref(erence)?|#)\s*$")
IS_PHONE = re.compile(r"\b(phone|mobile|tel|contact no\.?)\b")
IS_PINCODE = re.compile(r"\b(pin\s?code|pincode|postal code|zip)\b")
LOOKALIKE_BETWEEN_DIGITS = re.compile(r"(?<=[0-9])[OoIl](?=[0-9])")


def gstin_check_digit_ok(gstin: str) -> bool:
    """The 15th character of a GSTIN is a mod-36 check digit over the first 14."""
    total = 0
    for i, ch in enumerate(gstin[:14]):
        value = GSTIN_CHARSET.index(ch) * (2 if i % 2 else 1)
        total += value // 36 + value % 36
    return GSTIN_CHARSET[(36 - total % 36) % 36] == gstin[14]


def flat(text: str) -> str:
    """Normalise text for the "is this value in the sheet" check (spacing, commas, dashes)."""
    text = str(text).replace("–", "-").replace("—", "-").replace("⁠", "").replace("​", "")
    return re.sub(r"[\s,;|]+", "", text).lower()


def validate_header_fields(header_fields: list, sheet_text: str) -> list:
    """Rule checks on the model's header fields. Also tidies PAN/GSTIN values in place."""
    warnings = []
    source = flat(sheet_text)
    pans, gstins, po_numbers = [], [], []

    for field in header_fields:
        name = str(field.get("field", ""))
        lname = name.lower()
        value = str(field.get("value", "")).strip()

        # Catches values the model reworded, merged or made up.
        if value and flat(value) not in source:
            warnings.append(f"'{name}' = '{value}' not found verbatim in the sheet, "
                            f"the model may have reworded or merged it.")

        # A code/number/date field holding a long block of text usually means a blank label
        # picked up the next block (Scootsy's "Reference PO Code").
        is_short_field = (IS_ID_FIELD.search(lname) or IS_PO_NUMBER.search(lname) or "date" in lname) \
            and not (IS_GSTIN.search(lname) or IS_PAN.search(lname) or IS_PHONE.search(lname))
        if is_short_field and (len(value) > 40 or len(value.split()) > 6):
            warnings.append(f"'{name}' should be a short code/number/date but holds a long text block "
                            f"('{value[:60]}...'), probably text from another block attached to a blank label.")

        if IS_PAN.search(lname):
            value = value.replace(" ", "").upper()
            # strip a leftover label, e.g. "PAN:AAICP2804J"
            embedded = re.search(r"(?<![A-Z0-9])[A-Z]{5}[0-9]{4}[A-Z](?![A-Z0-9])", value)
            if embedded and embedded.group(0) != value:
                value = embedded.group(0)
            field["value"] = value
            pans.append((name, value))
            if not PAN_RE.match(value):
                warnings.append(f"PAN '{name}' = '{value}' doesn't match the AAAAA9999A pattern.")
            continue

        if IS_GSTIN.search(lname):
            value = value.replace(" ", "").upper()
            # strip a leftover label, e.g. BigBasket's "GSTIN NO:GST: 19AACCI2053A1Z4"
            embedded = re.search(r"(?<![A-Z0-9])[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z](?![A-Z0-9])", value)
            if embedded and embedded.group(0) != value:
                value = embedded.group(0)
            field["value"] = value
            if not GSTIN_RE.match(value):
                warnings.append(f"GSTIN '{name}' = '{value}' doesn't match the 15-character GSTIN format.")
            elif not gstin_check_digit_ok(value):
                warnings.append(f"GSTIN '{name}' = '{value}' fails its check digit, a character is wrong.")
            else:
                gstins.append((name, value))
            continue

        if IS_PO_NUMBER.search(lname) and "date" not in lname:
            po_numbers.append((name, value))
            if re.search(r"\s", value):
                warnings.append(f"Document number '{name}' = '{value}' contains spaces; another field "
                                f"may have been merged into it.")

        if IS_PHONE.search(lname):
            digits = re.sub(r"\D", "", value)
            if re.search(r"[A-Za-z]", value) or not 10 <= len(digits) <= 12:
                warnings.append(f"Phone '{name}' = '{value}' isn't a 10-12 digit phone number.")
            continue

        if IS_PINCODE.search(lname) and not re.fullmatch(r"[1-9][0-9]{5}", value.replace(" ", "")):
            warnings.append(f"Pincode '{name}' = '{value}' isn't a 6-digit pincode.")
            continue

        if (IS_ID_FIELD.search(lname) or IS_PO_NUMBER.search(lname)) and LOOKALIKE_BETWEEN_DIGITS.search(value):
            warnings.append(f"ID '{name}' = '{value}' has a letter O/I/l between digits, "
                            f"probably meant 0 or 1.")

    if not po_numbers:
        warnings.append("No document number (PO / STR / STN / STO number) was found in the header fields.")

    # Characters 3-12 of a GSTIN are the owner's PAN. A near match (1-3 chars off) with a
    # PAN on the sheet means one of the two is misread.
    for g_name, gstin in gstins:
        embedded = gstin[2:12]
        if any(embedded == pan for _, pan in pans):
            continue
        for p_name, pan in pans:
            if len(pan) == 10:
                diff = sum(a != b for a, b in zip(embedded, pan))
                if 1 <= diff <= 3:
                    warnings.append(f"GSTIN '{g_name}' embeds PAN '{embedded}', but '{p_name}' is "
                                    f"'{pan}' ({diff} character(s) different), one of them is wrong.")

    # Every GSTIN/PAN/CIN should be its own field, not just buried in an address block.
    extracted = {re.sub(r"\s", "", str(f.get("value", ""))).upper() for f in header_fields}
    for kind, pattern in IDS_IN_TEXT.items():
        for found in sorted(set(pattern.findall(sheet_text))):
            if found not in extracted:
                warnings.append(f"{kind} '{found}' is in the sheet but wasn't extracted as its own field.")

    return warnings


# --- output ---

def _unique(names: list) -> list:
    """Repeated column names get a number ("Qty", "Qty 2") so no value is lost in a dict."""
    seen, out = {}, []
    for name in names:
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f"{name} {seen[name]}")
    return out


def build_output(header_fields: list, totals: list, currency: str,
                 line_columns: list, line_rows: list, warnings: list) -> dict:
    """The raw extraction: header and totals keyed by their printed names, one dict per line item."""
    columns = _unique(line_columns)
    totals_out = {"Currency": currency} if currency else {}
    totals_out.update({t["field"]: t["value"] for t in totals})
    return {
        "header": {f.get("field", ""): f.get("value", "") for f in header_fields},
        "totals": totals_out,
        "line_item_columns": columns,
        "line_items": [dict(zip(columns, row)) for row in line_rows],
        "warnings": warnings,
    }


class PipelineError(RuntimeError):
    """Raised when a run fails (bad file, AWS or model error)."""


def process(sheet_path: str, clients: dict = None) -> dict:
    """
    Run the pipeline on one PO spreadsheet and return the extracted data as a dict.
    Raises PipelineError on failure. Called from run_po.py and main().
    """
    run_started = time.time()
    po_path = sheet_path
    clients = clients or make_clients()

    # Line items first: it's free, so a bad file fails before we pay for the model call.
    try:
        print(f"Reading PO from: {po_path}")
        fmt, rows = read_sheet_rows(po_path, read_file_bytes(clients["s3"], po_path))
        print(f"  [Input] format {fmt.upper()}, {len(rows)} row(s)")
        header_idx, serial_col, qty_col = find_table_header(rows)
        items = read_line_items(rows, header_idx, serial_col, qty_col)
        how = "serial column" if serial_col is not None else "Quantity column"
        print(f"  [Line items] header on row {header_idx + 1} (by {how}): "
              f"{len(items['rows'])} line item(s), {len(items['columns'])} column(s), "
              f"{len(items['totals'])} total field(s)\n")
    except NoCredentialsError as e:
        raise PipelineError("No AWS credentials found. Run `aws configure` first.") from e
    except ClientError as e:
        error_code = e.response.get("Error", {}).get("Code", "Unknown")
        raise PipelineError(f"{error_code} - {e}") from e
    except (OSError, ValueError, RuntimeError) as e:
        raise PipelineError(str(e)) from e

    sheet_text = header_section_text(rows, header_idx)
    if sheet_text:
        print(f"Calling {BEDROCK_MODEL_ID} for header fields "
              f"({len(sheet_text)} chars from rows 1-{header_idx})...")
        try:
            header_result = call_model_for_header(clients["bedrock"], sheet_text)
        except Exception as e:
            raise PipelineError(f"calling Bedrock (header fields): {e}") from e
        header_warnings = validate_header_fields(header_result["header_fields"], sheet_text)
    else:
        print("Nothing above the line-items table, skipping the model call.")
        header_result = {"header_fields": [], "cost": 0.0}
        header_warnings = []

    header_fields = header_result["header_fields"]
    line_columns, line_rows, totals = items["columns"], items["rows"], items["totals"]

    arithmetic_warnings, arithmetic_summary = validate_line_items(line_columns, line_rows)
    totals_warnings, totals_summary = reconcile_totals(items["header"], items["raw_rows"], totals)
    print(f"  [Line items] {arithmetic_summary}")
    print(f"  [Totals] {totals_summary}")

    print(f"\nExtracted {len(header_fields)} header field(s), {len(totals)} total(s) "
          f"and {len(line_rows)} line item(s) [{fmt.upper()} file].\n")

    all_warnings = header_warnings + totals_warnings + arithmetic_warnings
    if all_warnings:
        print(f"Validation: {len(all_warnings)} issue(s) to review:")
        for w in all_warnings:
            print(f"  - {w}")
    else:
        print("Validation: no issues.")

    print(f"\nCost for this document:")
    print(f"  {fmt.upper()} file (line items + totals):  $0.0000")
    print(f"  Model, header (text):             ${header_result['cost']:.4f}")
    print(f"  TOTAL:                            ${header_result['cost']:.4f}\n")

    print(f"Took {time.time() - run_started:.1f}s")
    return build_output(header_fields, totals, items["currency"], line_columns, line_rows, all_warnings)


def main():
    if len(sys.argv) != 2:
        print("Usage: python extract_excel_only_v2.py <PO file (.xls/.xlsx/.csv/.tsv): s3://... or local>")
        sys.exit(1)
    try:
        with contextlib.redirect_stdout(sys.stderr):      # log to stderr, JSON alone on stdout
            result = process(sys.argv[1])
    except PipelineError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)
    sys.stdout.reconfigure(encoding="utf-8")       # Windows defaults to cp1252 when redirected
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
